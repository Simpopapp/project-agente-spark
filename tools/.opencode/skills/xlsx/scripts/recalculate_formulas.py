"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import json
import os
import platform
import subprocess
import sys
from pathlib import Path

from office.run_libreoffice import get_soffice_env

from openpyxl import load_workbook

LO_MACRO_PATH_MACOS = "~/Library/Application Support/LibreOffice/4/user/basic/Standard"
LO_MACRO_PATH_LINUX = "~/.config/libreoffice/4/user/basic/Standard"
LO_MACRO_FILE = "Module1.xba"

RECALC_MACRO_SOURCE = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def gtimeout_available():
    try:
        subprocess.run(
            ["gtimeout", "--version"], capture_output=True, timeout=1, check=False
        )
        return True
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


def configure_lo_macro():
    macro_dir = os.path.expanduser(
        LO_MACRO_PATH_MACOS if platform.system() == "Darwin" else LO_MACRO_PATH_LINUX
    )
    macro_file = os.path.join(macro_dir, LO_MACRO_FILE)

    if (
        os.path.exists(macro_file)
        and "RecalculateAndSave" in Path(macro_file).read_text()
    ):
        return True

    if not os.path.exists(macro_dir):
        subprocess.run(
            ["soffice", "--headless", "--terminate_after_init"],
            capture_output=True,
            timeout=10,
            env=get_soffice_env(),
        )
        os.makedirs(macro_dir, exist_ok=True)

    try:
        Path(macro_file).write_text(RECALC_MACRO_SOURCE)
        return True
    except Exception:
        return False


def run_recalculation(filename, timeout=30):
    if not Path(filename).exists():
        return {"error": f"File {filename} does not exist"}

    abs_path = str(Path(filename).absolute())

    if not configure_lo_macro():
        return {"error": "Failed to setup LibreOffice macro"}

    lo_command = [
        "soffice",
        "--headless",
        "--norestore",
        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
        abs_path,
    ]

    if platform.system() == "Linux":
        lo_command = ["timeout", str(timeout)] + lo_command
    elif platform.system() == "Darwin" and gtimeout_available():
        lo_command = ["gtimeout", str(timeout)] + lo_command

    result = subprocess.run(lo_command, capture_output=True, text=True, env=get_soffice_env())

    if result.returncode != 0 and result.returncode != 124:  
        error_msg = result.stderr or "Unknown error during recalculation"
        if "Module1" in error_msg or "RecalculateAndSave" not in error_msg:
            return {"error": "LibreOffice macro not configured properly"}
        return {"error": error_msg}

    try:
        workbook = load_workbook(filename, data_only=True)

        error_patterns = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
        ]
        error_locations = {err: [] for err in error_patterns}
        error_count = 0

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        for err in error_patterns:
                            if err in cell.value:
                                location = f"{sheet_name}!{cell.coordinate}"
                                error_locations[err].append(location)
                                error_count += 1
                                break

        workbook.close()

        result = {
            "status": "success" if error_count == 0 else "errors_found",
            "total_errors": error_count,
            "error_summary": {},
        }

        for err_type, locations in error_locations.items():
            if locations:
                result["error_summary"][err_type] = {
                    "count": len(locations),
                    "locations": locations[:20],
                }

        formula_workbook = load_workbook(filename, data_only=False)
        total_formulas = 0
        for sheet_name in formula_workbook.sheetnames:
            worksheet = formula_workbook[sheet_name]
            for row in worksheet.iter_rows():
                for cell in row:
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):
                        total_formulas += 1
        formula_workbook.close()

        result["total_formulas"] = total_formulas

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalculate_formulas.py <excel_file> [timeout_seconds]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        sys.exit(1)

    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = run_recalculation(filename, timeout)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
